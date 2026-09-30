import io
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF


def sanitize_text(text):
    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u00a0": " "
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    return text.strip()


def format_docx(text, document_type, logo_path=None):
    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    if logo_path and Path(logo_path).exists():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run()
        run.add_picture(logo_path, width=Inches(1.2))

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title_run = title.add_run(document_type.upper())
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(16)

    cleaned_text = sanitize_text(text)

    for line in cleaned_text.splitlines():
        line = line.strip()

        if not line:
            document.add_paragraph()
            continue

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)

        run = paragraph.add_run(line)
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)

        if re.match(r"^(\d+\.|[A-Z][A-Z ]{3,})$", line):
            run.bold = True

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer_run = footer_paragraph.add_run(
        "LegalEase - Draft for Review"
    )
    footer_run.font.size = Pt(9)

    output = io.BytesIO()
    document.save(output)

    return output.getvalue()


class LegalEasePDF(FPDF):

    def __init__(self, logo_path=None):
        super().__init__()
        self.logo_path = logo_path

    def header(self):
        if self.logo_path and Path(self.logo_path).exists():
            try:
                self.image(
                    self.logo_path,
                    x=95,
                    y=8,
                    w=20
                )
                self.ln(15)
            except Exception:
                self.ln(5)
        else:
            self.ln(5)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", size=8)

        self.cell(
            0,
            10,
            "LegalEase - Draft for Review",
            align="C"
        )


def format_pdf(text, document_type, logo_path=None):
    pdf = LegalEasePDF(logo_path)

    pdf.set_auto_page_break(
        auto=True,
        margin=20
    )

    pdf.add_page()

    pdf.set_font("Helvetica", "B", 16)

    pdf.multi_cell(
        0,
        10,
        sanitize_text(document_type).upper(),
        align="C"
    )

    pdf.ln(4)

    cleaned_text = sanitize_text(text)

    for line in cleaned_text.splitlines():
        line = line.strip()

        if not line:
            pdf.ln(3)
            continue

        is_heading = bool(
            re.match(
                r"^(\d+\.|[A-Z][A-Z ]{3,})$",
                line
            )
        )

        if is_heading:
            pdf.set_font("Helvetica", "B", 11)
        else:
            pdf.set_font("Helvetica", "", 11)

        pdf.multi_cell(0, 7, line)
        pdf.ln(1)

    return bytes(pdf.output())


def format_html_preview(text):
    cleaned = sanitize_text(text)

    cleaned = (
        cleaned
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )

    cleaned = cleaned.replace("\n", "<br>")

    return cleaned